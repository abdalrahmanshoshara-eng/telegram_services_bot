"""Bounded Excel cleaning, splitting, and merging utilities."""

from __future__ import annotations

import os
import re
import zipfile
from copy import copy
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


class ExcelToolsError(ValueError):
    """Raised when an Excel workbook or requested operation is invalid."""


MAX_ARCHIVE_PARTS = 2_000
MAX_UNCOMPRESSED_BYTES = max(
    10, int(os.environ.get("MAX_EXCEL_UNCOMPRESSED_MB", "200"))
) * 1024 * 1024
MAX_WORKSHEETS = max(1, int(os.environ.get("MAX_EXCEL_WORKSHEETS", "100")))
MAX_WORKBOOK_CELLS = max(
    10_000, int(os.environ.get("MAX_EXCEL_CELLS", "1000000"))
)
MAX_SPLIT_GROUPS = max(
    2, int(os.environ.get("MAX_EXCEL_SPLIT_GROUPS", "50"))
)


def _is_empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _validate_xlsx_container(path: Path) -> None:
    if path.suffix.lower() != ".xlsx":
        raise ExcelToolsError("هذه الخدمة تدعم ملفات XLSX فقط.")
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_PARTS:
                raise ExcelToolsError("يحتوي ملف Excel على عدد مفرط من الأجزاء.")
            if sum(member.file_size for member in members) > MAX_UNCOMPRESSED_BYTES:
                raise ExcelToolsError("حجم ملف Excel بعد فك الضغط يتجاوز الحد المسموح.")
            names = {member.filename.lower() for member in members}
            if "encryptedpackage" in names or "encryptioninfo" in names:
                raise ExcelToolsError("ملفات Excel المشفرة غير مدعومة.")
            if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
                raise ExcelToolsError("محتوى الملف لا يطابق صيغة XLSX.")
    except ExcelToolsError:
        raise
    except (OSError, zipfile.BadZipFile) as error:
        raise ExcelToolsError("ملف Excel تالف أو غير صالح.") from error


def _load(path: Path):
    _validate_xlsx_container(path)
    try:
        workbook = load_workbook(path, data_only=False, keep_links=False)
    except Exception as error:
        raise ExcelToolsError("تعذر فتح ملف Excel. تأكد من سلامته.") from error
    if len(workbook.worksheets) > MAX_WORKSHEETS:
        workbook.close()
        raise ExcelToolsError(f"عدد أوراق العمل يتجاوز الحد ({MAX_WORKSHEETS}).")
    estimated_cells = sum(
        max(1, sheet.max_row) * max(1, sheet.max_column)
        for sheet in workbook.worksheets
    )
    if estimated_cells > MAX_WORKBOOK_CELLS:
        workbook.close()
        raise ExcelToolsError(
            f"حجم بيانات Excel يتجاوز الحد ({MAX_WORKBOOK_CELLS:,} خلية)."
        )
    return workbook


def _row_values(sheet: Worksheet, row_index: int) -> list:
    return [sheet.cell(row_index, column).value for column in range(1, sheet.max_column + 1)]


def _row_is_empty(sheet: Worksheet, row_index: int) -> bool:
    return all(_is_empty(value) for value in _row_values(sheet, row_index))


def _first_content_row(sheet: Worksheet) -> int:
    for row_index in range(1, sheet.max_row + 1):
        if not _row_is_empty(sheet, row_index):
            return row_index
    raise ExcelToolsError(f"ورقة العمل «{sheet.title}» فارغة.")


def _canonical_value(value):
    if isinstance(value, str):
        return value.strip()
    try:
        hash(value)
        return value
    except TypeError:
        return repr(value)


def _copy_cell(source, target) -> None:
    target.value = source.value
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)
    if source.comment:
        target.comment = copy(source.comment)


def _copy_row(source: Worksheet, target: Worksheet, source_row: int, target_row: int) -> None:
    for column in range(1, source.max_column + 1):
        _copy_cell(source.cell(source_row, column), target.cell(target_row, column))
    source_dimension = source.row_dimensions[source_row]
    if source_dimension.height is not None:
        target.row_dimensions[target_row].height = source_dimension.height


def _copy_dimensions(source: Worksheet, target: Worksheet) -> None:
    for index in range(1, source.max_column + 1):
        letter = get_column_letter(index)
        source_dimension = source.column_dimensions[letter]
        if source_dimension.width is not None:
            target.column_dimensions[letter].width = source_dimension.width
        target.column_dimensions[letter].hidden = source_dimension.hidden
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines


def _copy_sheet(source: Worksheet, target: Worksheet) -> None:
    _copy_dimensions(source, target)
    for row_index in range(1, source.max_row + 1):
        _copy_row(source, target, row_index, row_index)
    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))
    target.sheet_properties.pageSetUpPr = copy(source.sheet_properties.pageSetUpPr)
    target.page_setup = copy(source.page_setup)
    target.page_margins = copy(source.page_margins)
    target.print_options = copy(source.print_options)


def clean_excel_workbook(source: Path, destination: Path) -> dict[str, int]:
    workbook = _load(source)
    removed_blank = 0
    removed_duplicates = 0
    trimmed_cells = 0
    try:
        for sheet in workbook.worksheets:
            try:
                header_row = _first_content_row(sheet)
            except ExcelToolsError:
                continue

            seen: set[tuple] = set()
            rows_to_delete: list[int] = []
            for row_index in range(1, sheet.max_row + 1):
                values = _row_values(sheet, row_index)
                if all(_is_empty(value) for value in values):
                    rows_to_delete.append(row_index)
                    removed_blank += 1
                    continue

                for cell in sheet[row_index]:
                    if isinstance(cell.value, str):
                        cleaned = cell.value.strip()
                        if cleaned != cell.value:
                            cell.value = cleaned
                            trimmed_cells += 1

                if row_index <= header_row:
                    continue
                key = tuple(_canonical_value(value) for value in _row_values(sheet, row_index))
                if key in seen:
                    rows_to_delete.append(row_index)
                    removed_duplicates += 1
                else:
                    seen.add(key)

            for row_index in reversed(rows_to_delete):
                sheet.delete_rows(row_index)

        workbook.save(destination)
    finally:
        workbook.close()
    return {
        "blank_rows": removed_blank,
        "duplicate_rows": removed_duplicates,
        "trimmed_cells": trimmed_cells,
    }


def _normalize_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _resolve_split_column(sheet: Worksheet, selector: str) -> tuple[int, int, str]:
    selector = selector.strip()
    if not selector:
        raise ExcelToolsError("أرسل اسم العمود أو حرفه.")
    if re.fullmatch(r"[A-Za-z]{1,3}", selector):
        column = column_index_from_string(selector.upper())
        if column > sheet.max_column:
            raise ExcelToolsError("حرف العمود خارج نطاق البيانات.")
        header_row = _first_content_row(sheet)
        label = str(sheet.cell(header_row, column).value or selector.upper())
        return header_row, column, label

    expected = _normalize_header(selector)
    for row_index in range(1, min(sheet.max_row, 20) + 1):
        for column in range(1, sheet.max_column + 1):
            if _normalize_header(sheet.cell(row_index, column).value) == expected:
                return row_index, column, str(sheet.cell(row_index, column).value)
    raise ExcelToolsError(f"لم يتم العثور على العمود «{selector}» في أول 20 صفًا.")


def _safe_filename(value: str, fallback: str) -> str:
    cleaned = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", value).strip(" ._")
    cleaned = re.sub(r"\s+", " ", cleaned)
    return (cleaned or fallback)[:70]


def split_excel_workbook(
    source: Path,
    destination_zip: Path,
    column_selector: str,
) -> dict[str, int]:
    workbook = _load(source)
    try:
        sheet = workbook.active
        header_row, column_index, header_label = _resolve_split_column(
            sheet, column_selector
        )
        groups: dict[str, list[int]] = {}
        skipped = 0
        for row_index in range(header_row + 1, sheet.max_row + 1):
            if _row_is_empty(sheet, row_index):
                continue
            raw_value = sheet.cell(row_index, column_index).value
            if _is_empty(raw_value):
                skipped += 1
                continue
            label = str(raw_value).strip()
            groups.setdefault(label, []).append(row_index)

        if not groups:
            raise ExcelToolsError("لا توجد قيم قابلة للتقسيم في العمود المحدد.")
        if len(groups) > MAX_SPLIT_GROUPS:
            raise ExcelToolsError(
                f"عدد القيم المختلفة يتجاوز الحد ({MAX_SPLIT_GROUPS} ملفًا)."
            )

        used_names: set[str] = set()
        with zipfile.ZipFile(
            destination_zip, "w", zipfile.ZIP_DEFLATED, compresslevel=6
        ) as archive:
            for index, (label, row_indexes) in enumerate(groups.items(), start=1):
                output_book = Workbook()
                output_sheet = output_book.active
                output_sheet.title = sheet.title[:31]
                _copy_dimensions(sheet, output_sheet)
                target_row = 1
                for source_row in range(1, header_row + 1):
                    _copy_row(sheet, output_sheet, source_row, target_row)
                    target_row += 1
                for source_row in row_indexes:
                    _copy_row(sheet, output_sheet, source_row, target_row)
                    target_row += 1

                base_name = _safe_filename(label, f"group-{index:02d}")
                candidate = base_name
                suffix = 2
                while candidate.casefold() in used_names:
                    candidate = f"{base_name[:60]}-{suffix}"
                    suffix += 1
                used_names.add(candidate.casefold())
                output_path = destination_zip.parent / f"{candidate}.xlsx"
                output_book.save(output_path)
                output_book.close()
                archive.write(output_path, arcname=f"files/{candidate}.xlsx")
                output_path.unlink()

            summary = (
                "تقرير تقسيم Excel\n"
                f"الورقة: {sheet.title}\n"
                f"العمود: {header_label}\n"
                f"عدد الملفات: {len(groups)}\n"
                f"عدد صفوف البيانات: {sum(len(rows) for rows in groups.values())}\n"
                f"صفوف بلا قيمة في عمود التقسيم: {skipped}\n"
            )
            archive.writestr("summary.txt", summary.encode("utf-8-sig"))
        return {
            "groups": len(groups),
            "rows": sum(len(rows) for rows in groups.values()),
            "skipped": skipped,
        }
    finally:
        workbook.close()


def _unique_sheet_title(workbook: Workbook, desired: str) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", desired).strip() or "Sheet"
    cleaned = cleaned[:31]
    existing = {sheet.title.casefold() for sheet in workbook.worksheets}
    if cleaned.casefold() not in existing:
        return cleaned
    counter = 2
    while True:
        suffix = f"-{counter}"
        candidate = f"{cleaned[:31 - len(suffix)]}{suffix}"
        if candidate.casefold() not in existing:
            return candidate
        counter += 1


def merge_excel_workbooks(
    sources: Iterable[tuple[str, Path]],
    destination: Path,
    mode: str,
) -> dict[str, int]:
    source_list = list(sources)
    if len(source_list) < 2:
        raise ExcelToolsError("أرسل ملفي Excel على الأقل.")
    if mode not in {"rows", "sheets"}:
        raise ExcelToolsError("طريقة دمج Excel غير معروفة.")

    output_book = Workbook()
    output_book.remove(output_book.active)
    opened = []
    merged_rows = 0
    merged_sheets = 0
    try:
        if mode == "sheets":
            for filename, path in source_list:
                workbook = _load(path)
                opened.append(workbook)
                for source_sheet in workbook.worksheets:
                    if len(output_book.worksheets) >= MAX_WORKSHEETS:
                        raise ExcelToolsError(
                            f"مجموع أوراق العمل يتجاوز الحد ({MAX_WORKSHEETS})."
                        )
                    stem = Path(filename).stem
                    title = _unique_sheet_title(
                        output_book, f"{stem}-{source_sheet.title}"
                    )
                    target_sheet = output_book.create_sheet(title)
                    _copy_sheet(source_sheet, target_sheet)
                    merged_sheets += 1
        else:
            target_sheet = output_book.create_sheet("البيانات المدمجة")
            expected_headers: list[str] | None = None
            target_row = 1
            for filename, path in source_list:
                workbook = _load(path)
                opened.append(workbook)
                source_sheet = workbook.active
                header_row = _first_content_row(source_sheet)
                headers = [
                    _normalize_header(value)
                    for value in _row_values(source_sheet, header_row)
                ]
                while headers and not headers[-1]:
                    headers.pop()
                if not headers or any(not value for value in headers):
                    raise ExcelToolsError(
                        f"صف العناوين في الملف «{filename}» يحتوي خلايا فارغة."
                    )
                if expected_headers is None:
                    expected_headers = headers
                    _copy_dimensions(source_sheet, target_sheet)
                    _copy_row(source_sheet, target_sheet, header_row, target_row)
                    target_row += 1
                elif headers != expected_headers:
                    raise ExcelToolsError(
                        f"عناوين الأعمدة في الملف «{filename}» لا تطابق الملف الأول."
                    )

                for row_index in range(header_row + 1, source_sheet.max_row + 1):
                    if _row_is_empty(source_sheet, row_index):
                        continue
                    _copy_row(source_sheet, target_sheet, row_index, target_row)
                    target_row += 1
                    merged_rows += 1
            merged_sheets = 1

        output_book.save(destination)
    finally:
        output_book.close()
        for workbook in opened:
            workbook.close()
    return {
        "files": len(source_list),
        "rows": merged_rows,
        "sheets": merged_sheets,
    }
