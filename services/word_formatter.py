from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table


def _set_repeat_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    if row_properties.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        row_properties.append(header)


def _prevent_row_split(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    if row_properties.find(qn("w:cantSplit")) is None:
        row_properties.append(OxmlElement("w:cantSplit"))


def _iter_tables(container):
    for table in container.tables:
        yield table
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_tables(cell)


def _keep_small_table_together(table: Table) -> None:
    """Ask Word to move a table to the next page when it fits there.

    Word may still split tables taller than a full page; each row remains protected
    separately by ``w:cantSplit``.
    """
    last_row_index = len(table.rows) - 1
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = row_index < last_row_index


def format_word_tables(source: Path, destination: Path) -> None:
    document = Document(str(source))
    for table in _iter_tables(document):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        if table.rows:
            _set_repeat_header(table.rows[0])

        _keep_small_table_together(table)

        for row in table.rows:
            _prevent_row_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.keep_together = True

    document.save(str(destination))
