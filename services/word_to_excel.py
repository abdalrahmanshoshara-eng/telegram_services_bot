"""Collect the text of several DOCX files into one Excel sheet.

Each document becomes one row: its original file name, then its full text in
reading order - paragraphs and table rows interleaved exactly as they appear in
the body, not tables appended at the end.

An Excel cell holds at most 32,767 characters. Rather than silently cutting a
long document short, the remainder continues in the next column(s) of the same
row, and those continuation columns get their own headers.
"""

from __future__ import annotations

import os
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class WordToExcelError(ValueError):
    """Raised when a document cannot be read; the message is shown to the user."""


EXCEL_CELL_LIMIT = 32_767
MAX_ARCHIVE_PARTS = 2_000
MAX_UNCOMPRESSED_BYTES = max(
    10, int(os.environ.get("MAX_DOCX_UNCOMPRESSED_MB", "200"))
) * 1024 * 1024

NAME_HEADER = "اسم الملف"
CONTENT_HEADER = "المحتوى"

# Control characters that are not allowed in XLSX cell XML (tab, LF, CR are).
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_BLANK_LINES = re.compile(r"\n{3,}")


def _validate_docx_container(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_PARTS:
                raise WordToExcelError("يحتوي ملف Word على عدد مفرط من الأجزاء.")
            if sum(member.file_size for member in members) > MAX_UNCOMPRESSED_BYTES:
                raise WordToExcelError("حجم ملف Word بعد فك الضغط يتجاوز الحد المسموح.")
            names = {member.filename.lower() for member in members}
            if "encryptedpackage" in names or "encryptioninfo" in names:
                raise WordToExcelError("ملفات Word المحمية بكلمة مرور غير مدعومة.")
            if "word/document.xml" not in names:
                raise WordToExcelError("محتوى الملف لا يطابق صيغة DOCX.")
    except WordToExcelError:
        raise
    except (OSError, zipfile.BadZipFile) as error:
        raise WordToExcelError("ملف Word تالف أو غير صالح.") from error


def _table_lines(table: Table) -> list[str]:
    lines = []
    for row in table.rows:
        cells = []
        previous = None
        for cell in row.cells:
            # Horizontally merged cells are returned once per grid column.
            if previous is not None and cell._tc is previous:
                continue
            previous = cell._tc
            cells.append(" ".join(cell.text.split()))
        line = " | ".join(cells).strip(" |")
        if line:
            lines.append(line)
    return lines


def extract_docx_text(path: Path) -> str:
    """Return the body text of a DOCX in reading order."""
    _validate_docx_container(path)
    try:
        document = Document(str(path))
    except Exception as error:
        raise WordToExcelError("تعذر فتح ملف Word. تأكد من سلامته.") from error

    lines: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            lines.append(Paragraph(child, document).text)
        elif child.tag == qn("w:tbl"):
            lines.extend(_table_lines(Table(child, document)))
    text = "\n".join(line.rstrip() for line in lines)
    return _BLANK_LINES.sub("\n\n", _ILLEGAL_XML.sub("", text)).strip()


def _chunks(text: str) -> list[str]:
    if not text:
        return [""]
    return [
        text[start : start + EXCEL_CELL_LIMIT]
        for start in range(0, len(text), EXCEL_CELL_LIMIT)
    ]


def _set_text(cell, value: str) -> None:
    cell.value = value
    # A document line such as "=SUM(A1)" must stay text, never become a formula.
    cell.data_type = "s"


def word_files_to_excel(
    sources: list[tuple[str, Path]], destination: Path
) -> dict[str, int]:
    """Write one row per (display_name, path) and return a short summary."""
    if not sources:
        raise WordToExcelError("أرسل ملف Word واحدًا على الأقل.")

    rows = []
    for name, path in sources:
        try:
            text = extract_docx_text(path)
        except WordToExcelError as error:
            raise WordToExcelError(f"{name}: {error}") from error
        rows.append((_ILLEGAL_XML.sub("", name), _chunks(text), len(text)))

    content_columns = max(len(chunks) for _, chunks, _ in rows)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "محتوى الملفات"
    sheet.sheet_view.rightToLeft = True

    headers = [NAME_HEADER, CONTENT_HEADER] + [
        f"{CONTENT_HEADER} (تابع {index})" for index in range(2, content_columns + 1)
    ]
    header_fill = PatternFill("solid", fgColor="DDE7E2")
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column)
        _set_text(cell, header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    wrap = Alignment(wrap_text=True, vertical="top")
    for row_index, (name, chunks, _) in enumerate(rows, start=2):
        name_cell = sheet.cell(row=row_index, column=1)
        _set_text(name_cell, name)
        name_cell.alignment = wrap
        for offset, chunk in enumerate(chunks):
            cell = sheet.cell(row=row_index, column=2 + offset)
            _set_text(cell, chunk)
            cell.alignment = wrap

    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 35
    for column in range(2, content_columns + 2):
        sheet.column_dimensions[get_column_letter(column)].width = 100

    workbook.save(destination)
    return {
        "files": len(rows),
        "characters": sum(length for _, _, length in rows),
        "empty": sum(1 for _, _, length in rows if length == 0),
        "split": sum(1 for _, chunks, _ in rows if len(chunks) > 1),
    }
