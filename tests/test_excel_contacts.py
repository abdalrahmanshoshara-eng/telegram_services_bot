import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from services.excel_contacts import ExcelContactsError, convert_excel_contacts
from services.excel_contacts.domain import normalize_phone, process_contact_rows


def make_xlsx(path, rows=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "جهات الاتصال"
    sheet.append(["الاسم الكامل", "رقم التواصل", "البريد الالكتروني"])
    for row in rows or [["مستخدم تجريبي", "0999123456", "user@example.com"]]:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def make_multisheet_xlsx(path):
    workbook = Workbook()
    unrelated = workbook.active
    unrelated.title = "الاحتياج"
    unrelated.append(["الرمز الوظيفي", "المسمى الوظيفي"])
    unrelated.append([1, "مهندس"])
    contacts = workbook.create_sheet("المرشحين")
    contacts.append(["رقم التقديم", "الاسم الكامل", "رقم التواصل", "البريد الالكتروني"])
    contacts.append([556368, "مستخدم تجريبي", 959070158, "user@example.com"])
    workbook.save(path)
    workbook.close()


def test_phone_normalization_and_classification_are_preserved():
    for value in [
        "0933123456",
        "933123456",
        "00963933123456",
        "+963 933 123 456",
        "٠٩٣٣١٢٣٤٥٦",
    ]:
        assert normalize_phone(value, "963") == "+963933123456"

    result = process_contact_rows(
        [
            {
                "الاسم الكامل": "أحمد الأول",
                "رقم التواصل": "0933123456",
                "البريد الالكتروني": "first@example.com",
            },
            {
                "الاسم الكامل": "أحمد المكرر",
                "رقم التواصل": "+963933123456",
                "البريد الالكتروني": "second@example.com",
            },
            {
                "الاسم الكامل": "ليلى",
                "رقم التواصل": "0944123456",
                "البريد الالكتروني": "invalid-email",
            },
        ]
    )
    assert result["summary"] == {
        "totalRows": 3,
        "validCount": 1,
        "duplicateCount": 1,
        "invalidCount": 1,
    }


def test_xlsx_is_converted_to_original_result_bundle(tmp_path):
    source = tmp_path / "contacts.xlsx"
    destination = tmp_path / "contacts-output.zip"
    vcf_destination = tmp_path / "contacts.vcf"
    make_xlsx(source, [["@اسم تجريبي", "0999123456", "user@example.com"]])

    summary = convert_excel_contacts(
        source,
        destination,
        "963",
        vcf_destination=vcf_destination,
    )

    assert summary["validCount"] == 1
    with zipfile.ZipFile(destination) as archive:
        assert set(archive.namelist()) == {
            "contacts.vcf",
            "clean_contacts.xlsx",
            "merged_duplicates.xlsx",
            "invalid_rows.xlsx",
            "summary.txt",
        }
        vcard = archive.read("contacts.vcf").decode("utf-8")
        assert "TEL;TYPE=CELL:+963999123456" in vcard
        assert "FN;CHARSET=UTF-8:@اسم تجريبي" in vcard
        report = load_workbook(
            io.BytesIO(archive.read("clean_contacts.xlsx")), data_only=False
        )
        assert report.active["A2"].value == "'@اسم تجريبي"
        assert report.active["B2"].value == "'+963999123456"
        report.close()
    assert vcf_destination.read_bytes() == vcard.encode("utf-8")


def test_contacts_sheet_is_found_when_it_is_not_the_first_sheet(tmp_path):
    source = tmp_path / "candidates.xlsx"
    destination = tmp_path / "contacts-output.zip"
    make_multisheet_xlsx(source)

    summary = convert_excel_contacts(source, destination, "963")

    assert summary == {
        "totalRows": 1,
        "validCount": 1,
        "duplicateCount": 0,
        "invalidCount": 0,
    }
    with zipfile.ZipFile(destination) as archive:
        assert "ورقة العمل المستخدمة: المرشحين" in archive.read(
            "summary.txt"
        ).decode("utf-8")
        assert "TEL;TYPE=CELL:+963959070158" in archive.read(
            "contacts.vcf"
        ).decode("utf-8")


def test_fake_excel_file_is_rejected(tmp_path):
    source = tmp_path / "contacts.xlsx"
    source.write_bytes(b"not an excel workbook")

    with pytest.raises(ExcelContactsError, match="لا يطابق صيغة Excel"):
        convert_excel_contacts(source, tmp_path / "output.zip", "963")
