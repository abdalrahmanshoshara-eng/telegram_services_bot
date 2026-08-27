from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from services.excel_tools import (
    ExcelToolsError,
    clean_excel_workbook,
    merge_excel_workbooks,
    split_excel_workbook,
)


def make_workbook(path, rows, sheet_name="Data"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 24
    workbook.save(path)
    workbook.close()


def test_clean_excel_removes_blank_duplicate_rows_and_trims_text(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "cleaned.xlsx"
    make_workbook(
        source,
        [
            [" الاسم ", "الإدارة"],
            [" أحمد ", "المالية"],
            [None, None],
            ["أحمد", "المالية"],
            ["سارة", " التقنية "],
        ],
    )

    summary = clean_excel_workbook(source, output)
    assert summary == {"blank_rows": 1, "duplicate_rows": 1, "trimmed_cells": 3}
    workbook = load_workbook(output)
    values = list(workbook.active.values)
    workbook.close()
    assert values == [
        ("الاسم", "الإدارة"),
        ("أحمد", "المالية"),
        ("سارة", "التقنية"),
    ]


def test_split_excel_creates_one_file_per_column_value(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "split.zip"
    make_workbook(
        source,
        [
            ["الاسم", "الإدارة"],
            ["أحمد", "المالية"],
            ["سارة", "التقنية"],
            ["ليلى", "المالية"],
            ["بلا إدارة", None],
        ],
    )

    summary = split_excel_workbook(source, output, "الإدارة")
    assert summary == {"groups": 2, "rows": 3, "skipped": 1}
    with ZipFile(output) as archive:
        names = archive.namelist()
        assert "summary.txt" in names
        assert len([name for name in names if name.endswith(".xlsx")]) == 2


def test_merge_excel_rows_and_sheets(tmp_path):
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    rows_output = tmp_path / "merged-rows.xlsx"
    sheets_output = tmp_path / "merged-sheets.xlsx"
    make_workbook(first, [["الاسم", "القسم"], ["أحمد", "أ"]], "First")
    make_workbook(second, [["الاسم", "القسم"], ["سارة", "ب"]], "Second")
    sources = [(first.name, first), (second.name, second)]

    summary = merge_excel_workbooks(sources, rows_output, "rows")
    assert summary == {"files": 2, "rows": 2, "sheets": 1}
    workbook = load_workbook(rows_output)
    assert list(workbook.active.values) == [
        ("الاسم", "القسم"),
        ("أحمد", "أ"),
        ("سارة", "ب"),
    ]
    workbook.close()

    summary = merge_excel_workbooks(sources, sheets_output, "sheets")
    assert summary == {"files": 2, "rows": 0, "sheets": 2}
    workbook = load_workbook(sheets_output)
    assert len(workbook.sheetnames) == 2
    assert list(workbook.worksheets[0].values)[1] == ("أحمد", "أ")
    assert list(workbook.worksheets[1].values)[1] == ("سارة", "ب")
    workbook.close()


def test_merge_rows_requires_matching_headers(tmp_path):
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    output = tmp_path / "merged.xlsx"
    make_workbook(first, [["الاسم"], ["أحمد"]])
    make_workbook(second, [["الهاتف"], ["123"]])

    with pytest.raises(ExcelToolsError, match="لا تطابق"):
        merge_excel_workbooks(
            [(first.name, first), (second.name, second)], output, "rows"
        )
