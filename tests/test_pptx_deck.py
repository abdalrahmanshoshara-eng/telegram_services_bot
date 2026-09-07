import zipfile
from pathlib import Path

import pytest
from docx import Document

from services.pptx_deck import (
    DEFAULT_SLIDE_COUNT,
    SLIDE_COUNT_CHOICES,
    ExtractionError,
    Outline,
    Slide,
    outline_from_dict,
    outline_preview,
    outline_to_dict,
    render_outline,
)
from services.pptx_deck import brand
from services.pptx_deck.extract import extract
from services.pptx_deck.rtl import estimate_lines, fix_bidi_digits, split_to_fit


def _sample_outline() -> Outline:
    return Outline(
        title="التقرير السنوي",
        subtitle="مراجعة الأداء",
        slides=[
            Slide(layout="cover", title="التقرير السنوي", subtitle="مراجعة الأداء"),
            Slide(layout="section", title="المحور الأول"),
            Slide(
                layout="bullets",
                title="أبرز المؤشرات",
                bullets=["ارتفاع نسبة الإنجاز", "تحسن زمن الاستجابة"],
            ),
            Slide(
                layout="two_col",
                title="مقارنة",
                columns=[
                    {"heading": "الأولى", "governorate": "دمشق", "bullets": ["بند"]},
                    {"heading": "الثانية", "governorate": "حلب", "bullets": ["بند"]},
                ],
            ),
            Slide(layout="stat", title="الإنجاز", stat={"value": "64%", "caption": "من الخطة"}),
            Slide(
                layout="chart",
                title="التوزع",
                chart={
                    "kind": "column",
                    "categories": ["الخدمات", "الرقمنة"],
                    "series": [{"name": "المخطط", "values": [28, 19]}],
                },
            ),
            Slide(layout="closing", title="شكراً لكم"),
        ],
    )


def _docx(path: Path, paragraphs: int = 8) -> Path:
    document = Document()
    document.add_heading("عنوان رئيسي", level=1)
    for index in range(paragraphs):
        document.add_paragraph(
            f"فقرة رقم {index} تحتوي على نص عربي كافٍ لتجاوز الحد الأدنى المطلوب "
            "لاستخراج المحتوى من المستند."
        )
    document.add_heading("محور فرعي", level=2)
    document.save(str(path))
    return path


# ---------------------------------------------------------------- assets ---
def test_identity_assets_are_committed():
    """The container filesystem is read-only, so assets cannot be fetched at runtime."""
    assert (brand.ASSETS / "fonts" / "HayyakumAllah-Regular.ttf").exists()
    assert brand.emblem("gold").exists()
    assert brand.governorate_icon("دمشق") is not None
    assert len(brand.GOVERNORATES) == 14


def test_governorate_lookup_tolerates_the_definite_article():
    assert brand.governorate_icon("اللاذقية") is not None
    assert brand.governorate_icon("لاذقية") is not None
    assert brand.governorate_icon("مدينة غير موجودة") is None


def test_flag_colours_are_not_reused_as_chart_colours():
    assert brand.FLAG_RED not in brand.CHART_SERIES
    assert brand.FLAG_GREEN not in brand.CHART_SERIES


# ------------------------------------------------------------- extraction ---
def test_docx_headings_seed_the_outline_hint(tmp_path):
    doc = extract(_docx(tmp_path / "sample.docx"))
    assert doc.page_count == 0
    assert len(doc.headings) == 2
    assert "عنوان رئيسي" in doc.outline_hint()


def test_short_document_is_rejected(tmp_path):
    document = Document()
    document.add_paragraph("قصير")
    path = tmp_path / "tiny.docx"
    document.save(str(path))
    with pytest.raises(ExtractionError):
        extract(path)


def test_unsupported_extension_is_rejected(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("محتوى", encoding="utf-8")
    with pytest.raises(ExtractionError):
        extract(path)


def test_extraction_errors_are_value_errors():
    """bot.py surfaces ValueError subclasses to the user verbatim."""
    assert issubclass(ExtractionError, ValueError)


# --------------------------------------------------------------- rtl bits ---
def test_digits_inside_arabic_get_a_right_to_left_mark():
    assert "‏" in fix_bidi_digits("بند رقم 3 من الخطة")
    assert fix_bidi_digits("64%") == "64%"          # standalone value untouched


def test_long_bullet_lists_split_across_slides():
    items = [f"بند تجريبي رقم {i} بمحتوى طويل نسبيًا" for i in range(20)]
    pages = split_to_fit(
        items, brand.CONTENT_W, brand.BODY_H, brand.SIZE_BODY, step=brand.LINE_STEP
    )
    assert len(pages) > 1
    assert sum(len(page) for page in pages) == len(items)


def test_line_estimation_grows_with_text_length():
    short = estimate_lines("نص قصير", brand.CONTENT_W, brand.SIZE_BODY)
    long = estimate_lines("نص طويل " * 60, brand.CONTENT_W, brand.SIZE_BODY)
    assert long > short


# ---------------------------------------------------------------- outline ---
def test_outline_survives_a_round_trip_through_job_state():
    original = _sample_outline()
    restored = outline_from_dict(outline_to_dict(original))
    assert restored.title == original.title
    assert [s.layout for s in restored.slides] == [s.layout for s in original.slides]
    assert restored.slides[4].stat == original.slides[4].stat


def test_preview_lists_every_slide():
    preview = outline_preview(_sample_outline())
    assert "أبرز المؤشرات" in preview
    assert "7 شريحة" in preview


def test_slide_count_choices_include_the_default():
    assert DEFAULT_SLIDE_COUNT in SLIDE_COUNT_CHOICES


# ----------------------------------------------------------------- render ---
def test_render_produces_a_deck_with_embedded_fonts(tmp_path):
    output = tmp_path / "deck.pptx"
    render_outline(_sample_outline(), output, organisation="وزارة")
    assert output.exists()

    with zipfile.ZipFile(output) as archive:
        names = archive.namelist()
        # One slide per outline entry.
        assert len([n for n in names if n.startswith("ppt/slides/slide")]) == 7
        # Fonts travel with the file - viewers do not need them installed.
        assert len([n for n in names if n.startswith("ppt/fonts/")]) == 4
        assert len([n for n in names if "charts/chart" in n]) >= 1
        presentation = archive.read("ppt/presentation.xml").decode("utf-8")
    assert 'embedTrueTypeFonts="1"' in presentation


def test_render_is_deterministic_in_slide_count(tmp_path):
    """Overflowing bullets add slides rather than spilling off the canvas."""
    outline = Outline(
        title="ت",
        subtitle="",
        slides=[
            Slide(layout="cover", title="غلاف"),
            Slide(
                layout="bullets",
                title="نقاط كثيرة",
                bullets=[f"بند رقم {i} بنص طويل نسبيًا للاختبار" for i in range(24)],
            ),
            Slide(layout="closing", title="شكراً"),
        ],
    )
    output = tmp_path / "overflow.pptx"
    render_outline(outline, output)
    with zipfile.ZipFile(output) as archive:
        slides = [n for n in archive.namelist() if n.startswith("ppt/slides/slide")]
    assert len(slides) > 3
