from zipfile import ZipFile

from docx import Document

from services.word_formatter import format_word_tables


def test_formats_table_and_adds_pagination_properties(tmp_path):
    source = tmp_path / "source.docx"
    destination = tmp_path / "formatted.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "العنوان"
    table.cell(1, 0).text = "المحتوى"
    document.save(source)

    format_word_tables(source, destination)

    assert destination.exists()
    with ZipFile(destination) as archive:
        xml = archive.read("word/document.xml")
    assert b"w:cantSplit" in xml
    assert b"w:tblHeader" in xml
    assert b'w:jc w:val="center"' in xml
