import pytest
from docx import Document
from openpyxl import load_workbook

from services.word_to_excel import (
    EXCEL_CELL_LIMIT,
    WordToExcelError,
    extract_docx_text,
    word_files_to_excel,
)


def _docx(path, *paragraphs, table=None, after=None):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    for text in after or []:
        document.add_paragraph(text)
    document.save(path)
    return path


def test_text_keeps_reading_order_with_tables(tmp_path):
    source = _docx(
        tmp_path / "a.docx",
        "مقدمة",
        table=[["الاسم", "العمر"], ["سامي", "30"]],
        after=["خاتمة"],
    )
    assert extract_docx_text(source) == "مقدمة\nالاسم | العمر\nسامي | 30\nخاتمة"


def test_one_row_per_file_with_name_and_content(tmp_path):
    first = _docx(tmp_path / "1.docx", "السطر الأول", "السطر الثاني")
    second = _docx(tmp_path / "2.docx", "=SUM(A1:A2)")
    output = tmp_path / "out.xlsx"

    summary = word_files_to_excel(
        [("تقرير أيلول.docx", first), ("معادلة.docx", second)], output
    )

    sheet = load_workbook(output).active
    assert [c.value for c in sheet[1]] == ["اسم الملف", "المحتوى"]
    assert [c.value for c in sheet[2]] == ["تقرير أيلول.docx", "السطر الأول\nالسطر الثاني"]
    assert sheet["B3"].value == "=SUM(A1:A2)"
    assert sheet["B3"].data_type == "s"
    assert summary == {"files": 2, "characters": 35, "empty": 0, "split": 0}


def test_long_content_continues_in_next_column(tmp_path):
    text = "ا" * (EXCEL_CELL_LIMIT + 10)
    source = _docx(tmp_path / "long.docx", text)
    output = tmp_path / "out.xlsx"

    summary = word_files_to_excel([("long.docx", source)], output)

    sheet = load_workbook(output).active
    assert sheet["C1"].value == "المحتوى (تابع 2)"
    assert len(sheet["B2"].value) == EXCEL_CELL_LIMIT
    assert sheet["B2"].value + sheet["C2"].value == text
    assert summary["split"] == 1


def test_invalid_file_names_the_culprit(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(WordToExcelError, match="bad.docx"):
        word_files_to_excel([("bad.docx", bad)], tmp_path / "out.xlsx")
